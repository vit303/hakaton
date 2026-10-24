"""Разбор документов: текст, таблицы и картинки."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from rag.config import CHUNK_OVERLAP, CHUNK_SIZE, SUPPORTED_EXTENSIONS, TABLE_CHUNK_SIZE


@dataclass
class Block:
    kind: str
    text: str
    page: int
    image: bytes | None = None
    image_name: str = ""


@dataclass
class ParsedDocument:
    source: str
    name: str
    blocks: list[Block] = field(default_factory=list)


def rows_to_markdown(rows: list[list[str]]) -> str:
    cleaned: list[list[str]] = []
    width = 0
    for row in rows:
        cells = [_cell(value) for value in row]
        if any(cells):
            cleaned.append(cells)
            width = max(width, len(cells))
    if not cleaned or width == 0:
        return ""
    normalized = [row + [""] * (width - len(row)) for row in cleaned]
    header = normalized[0]
    body = normalized[1:] or [[""] * width]
    lines = [
        _md_row(header),
        "| " + " | ".join("---" for _ in header) + " |",
    ]
    lines.extend(_md_row(row) for row in body)
    return "\n".join(lines)


def split_text(text: str, size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        return []
    if len(text) <= size:
        return [text]
    parts: list[str] = []
    start = 0
    while start < len(text):
        end = min(len(text), start + size)
        if end < len(text):
            cut = text.rfind("\n", start + size // 2, end)
            if cut == -1:
                cut = text.rfind(" ", start + size // 2, end)
            if cut != -1:
                end = cut
        piece = text[start:end].strip()
        if piece:
            parts.append(piece)
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return parts


def split_table(markdown: str, size: int = TABLE_CHUNK_SIZE) -> list[str]:
    markdown = markdown.strip()
    if not markdown:
        return []
    if len(markdown) <= size:
        return [markdown]
    lines = markdown.splitlines()
    header = "\n".join(lines[:2]) if len(lines) >= 2 else ""
    rows = lines[2:] if len(lines) >= 2 else lines
    chunks: list[str] = []
    current = header
    for row in rows:
        candidate = row if not current else current + "\n" + row
        if len(candidate) > size and current and current != header:
            chunks.append(current)
            current = header + "\n" + row if header else row
        else:
            current = candidate
    if current.strip():
        chunks.append(current)
    return chunks


def parse_file(path: Path) -> ParsedDocument:
    path = path.resolve()
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Формат не поддерживается: {path.name}")
    if not path.is_file():
        raise FileNotFoundError(path)
    parsers = {
        ".pdf": _parse_pdf,
        ".pptx": _parse_pptx,
        ".docx": _parse_docx,
        ".xlsx": _parse_xlsx,
        ".csv": _parse_csv,
        ".txt": _parse_text,
        ".md": _parse_text,
        ".png": _parse_image_file,
        ".jpg": _parse_image_file,
        ".jpeg": _parse_image_file,
        ".webp": _parse_image_file,
        ".bmp": _parse_image_file,
        ".tif": _parse_image_file,
        ".tiff": _parse_image_file,
    }
    blocks = parsers[suffix](path)
    return ParsedDocument(source=str(path), name=path.name, blocks=blocks)


def _cell(value: object) -> str:
    if value is None:
        return ""
    return str(value).replace("\n", " ").replace("|", "/").strip()


def _md_row(cells: list[str]) -> str:
    return "| " + " | ".join(cells) + " |"


def _parse_text(path: Path) -> list[Block]:
    raw = path.read_bytes()
    text = _decode(raw)
    return [Block("text", part, 1) for part in split_text(text)]


def _decode(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _parse_csv(path: Path) -> list[Block]:
    import pandas as pd

    last_error: Exception | None = None
    frame = None
    for encoding in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            frame = pd.read_csv(path, encoding=encoding)
            break
        except Exception as exc:  # noqa: BLE001 - пробуем следующую кодировку
            last_error = exc
    if frame is None:
        raise RuntimeError(f"Не удалось прочитать CSV {path.name}: {last_error}")
    rows = [list(frame.columns.astype(str))]
    rows.extend(frame.fillna("").astype(str).values.tolist())
    markdown = rows_to_markdown(rows)
    title = f"Таблица из файла {path.name}"
    return [Block("table", f"{title}\n{part}", 1) for part in split_table(markdown)]


def _parse_xlsx(path: Path) -> list[Block]:
    import openpyxl

    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    blocks: list[Block] = []
    try:
        for index, sheet in enumerate(workbook.worksheets, start=1):
            rows: list[list[str]] = []
            for row in sheet.iter_rows(values_only=True):
                cells = [_cell(value) for value in row]
                if any(cells):
                    rows.append(cells)
            markdown = rows_to_markdown(rows)
            if not markdown:
                continue
            title = f"Лист «{sheet.title}» файла {path.name}"
            for part in split_table(markdown):
                blocks.append(Block("table", f"{title}\n{part}", index))
    finally:
        workbook.close()
    return blocks


def _parse_docx(path: Path) -> list[Block]:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(str(path))
    blocks: list[Block] = []
    paragraphs: list[str] = []
    table_index = 0
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            text = Paragraph(child, document).text.strip()
            if text:
                paragraphs.append(text)
        elif child.tag == qn("w:tbl"):
            if paragraphs:
                for part in split_text("\n".join(paragraphs)):
                    blocks.append(Block("text", part, 1))
                paragraphs = []
            table_index += 1
            table = Table(child, document)
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            markdown = rows_to_markdown(rows)
            title = f"Таблица {table_index} файла {path.name}"
            for part in split_table(markdown):
                blocks.append(Block("table", f"{title}\n{part}", table_index))
    if paragraphs:
        for part in split_text("\n".join(paragraphs)):
            blocks.append(Block("text", part, 1))

    seen: set[str] = set()
    for rel in document.part.rels.values():
        if "image" not in rel.reltype:
            continue
        part = rel.target_part
        blob = getattr(part, "blob", None)
        if not blob or rel.rId in seen:
            continue
        seen.add(rel.rId)
        name = Path(str(getattr(part, "partname", "image"))).name
        blocks.append(Block("image", "", 1, image=blob, image_name=name))
    return blocks


def _parse_pptx(path: Path) -> list[Block]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(str(path))
    blocks: list[Block] = []
    for number, slide in enumerate(presentation.slides, start=1):
        texts: list[str] = []
        tables: list[str] = []
        images: list[tuple[bytes, str]] = []
        _walk_pptx_shapes(slide.shapes, texts, tables, images, MSO_SHAPE_TYPE)
        try:
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = (slide.notes_slide.notes_text_frame.text or "").strip()
                if notes:
                    texts.append("Заметки: " + notes)
        except Exception:
            pass
        page_text = "\n".join(piece for piece in texts if piece.strip())
        for part in split_text(page_text):
            blocks.append(Block("text", part, number))
        for index, markdown in enumerate(tables, start=1):
            title = f"Таблица {index} на слайде {number} файла {path.name}"
            for part in split_table(markdown):
                blocks.append(Block("table", f"{title}\n{part}", number))
        context = page_text[:500]
        for image, name in images:
            blocks.append(Block("image", context, number, image=image, image_name=name))
    return blocks


def _walk_pptx_shapes(shapes, texts: list[str], tables: list[str], images: list[tuple[bytes, str]], shape_type) -> None:
    for shape in shapes:
        if shape.shape_type == shape_type.GROUP:
            _walk_pptx_shapes(shape.shapes, texts, tables, images, shape_type)
            continue
        if getattr(shape, "has_table", False):
            rows = []
            for row in shape.table.rows:
                rows.append([cell.text for cell in row.cells])
            markdown = rows_to_markdown(rows)
            if markdown:
                tables.append(markdown)
            continue
        if shape.shape_type == shape_type.PICTURE:
            try:
                image = shape.image
                images.append((image.blob, f"slide.{image.ext or 'png'}"))
            except Exception:
                continue
            continue
        if getattr(shape, "has_text_frame", False):
            text = shape.text_frame.text.strip()
            if text:
                texts.append(text)


def _parse_pdf(path: Path) -> list[Block]:
    import pymupdf

    document = pymupdf.open(path)
    blocks: list[Block] = []
    seen_xrefs: set[int] = set()
    try:
        for index, page in enumerate(document, start=1):
            table_markdowns: list[str] = []
            try:
                finder = page.find_tables()
                for table in finder.tables:
                    extracted = table.extract() or []
                    markdown = rows_to_markdown(extracted)
                    if markdown:
                        table_markdowns.append(markdown)
            except Exception:
                table_markdowns = []
            text = page.get_text("text") or ""
            for part in split_text(text):
                blocks.append(Block("text", part, index))
            for number, markdown in enumerate(table_markdowns, start=1):
                title = f"Таблица {number} на странице {index} файла {path.name}"
                for part in split_table(markdown):
                    blocks.append(Block("table", f"{title}\n{part}", index))
            context = text.strip()[:500]
            for image in page.get_images(full=True):
                xref = image[0]
                if xref in seen_xrefs:
                    continue
                seen_xrefs.add(xref)
                try:
                    info = document.extract_image(xref)
                except Exception:
                    continue
                blob = info.get("image")
                if not blob:
                    continue
                ext = info.get("ext") or "png"
                blocks.append(Block("image", context, index, image=blob, image_name=f"page.{ext}"))
    finally:
        document.close()
    return blocks


def _parse_image_file(path: Path) -> list[Block]:
    return [Block("image", "", 1, image=path.read_bytes(), image_name=path.name)]
